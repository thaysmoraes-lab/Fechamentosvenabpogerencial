"""
VENA 2026 — Atualização Mensal Unificada
Planilhas: CMV | Estoque | FC Consolidado
"""

import re, shutil, os
import importlib
import pandas as pd
from openpyxl import load_workbook

# ─── Meses ───────────────────────────────────────────────────────────────────
MESES_DISPLAY = [
    "Janeiro","Fevereiro","Março","Abril","Maio","Junho",
    "Julho","Agosto","Setembro","Outubro","Novembro","Dezembro"
]
MES_PLANILHA = {
    "Janeiro":"JANEIRO","Fevereiro":"FEVEREIRO","Março":"MARÇO",
    "Abril":"ABRIL","Maio":"MAIO","Junho":"JUNHO",
    "Julho":"JULHO","Agosto":"AGOSTO","Setembro":"SETEMBRO",
    "Outubro":"OUTUBRO","Novembro":"NOVEMBRO","Dezembro":"DEZEMBRO"
}
MES_ABA_ESTOQUE = {
    "Janeiro":"JAN","Fevereiro":"FEV","Março":"MAR","Abril":"ABR",
    "Maio":"MAI","Junho":"JUN","Julho":"JUL","Agosto":"AGO",
    "Setembro":"SET","Outubro":"OUT","Novembro":"NOV","Dezembro":"DEZ"
}
MES_FC_PREFIX = {
    "Janeiro":"JAN","Fevereiro":"FEV","Março":"MAR","Abril":"ABR",
    "Maio":"MAI","Junho":"JUN","Julho":"JUL","Agosto":"AGO",
    "Setembro":"SET","Outubro":"OUT","Novembro":"NOV","Dezembro":"DEZ"
}

# Mapeamento Fornecedor/Grupo → linha Excel (Real 2022)
GRUPO_LINHA = {
    "1.Calamo":             20,
    "2.OBF":                21,
    "3.Antilhas":           22,
    "4.Outros Fornecedores":23,
    "5.Impostos":           24,
    "6.Folha":              25,
}
LINHA_NOME = {
    20:"Mercadorias",21:"Royalties",22:"Embalagens",
    23:"Outros Fornecedores",24:"Obrig. Tributárias",
    25:"Salários",26:"Outras Obrigações"
}

# Colunas CMV
COL_FAT       = {"ReceitaBruta":2,"Descontos":3,"ReceitaLiquida":4,"Trocas":5,"LiquidaMenosTrocas":6}
COL_CMV_VD    = {"CMV":9,"TRF":14}
COL_CMV_LOJAS = {"CMV":10,"TRF":15}
COL_VD        = {"ValorTabela":2,"ValorPraticado":3,"ValorPedido":4,"CCResidual":5}

# ─── Helpers ─────────────────────────────────────────────────────────────────
def br_float(val):
    if pd.isna(val): return None
    s = str(val).strip().replace(".","").replace(",",".")
    try: return float(s)
    except: return None

def store_code(s):
    m = re.match(r"^(\d+)\s*[-–]", str(s).strip())
    return int(m.group(1)) if m else None

def build_month_map(ws, mes):
    """
    Mapeia código de loja → linha Excel do mês.
    Suporta dois layouts:
      - Alcina: col A tem código+nome, col A tem mês (mesma coluna)
      - Righi:  col B tem código+nome, col C tem mês (colunas diferentes)
    Estratégia: varre todas as células da linha procurando código de loja
    E todas as células procurando o mês. Mantém o último código visto.
    """
    mes_upper = mes.strip().upper()
    result, current_code = {}, None
    for row in ws.iter_rows():
        found_mes  = False
        found_code = None
        for cell in row:
            if cell.value is None: continue
            cs = str(cell.value).strip()
            if not cs: continue
            # Reset em blocos de total/consolidado
            if "CONSOLIDADO" in cs.upper() or "TOTAL GERAL" in cs.upper():
                current_code = None
                break
            # Verifica se é código de loja
            code = store_code(cs)
            if code:
                found_code = code
            # Verifica se é o mês procurado
            if cs.upper() == mes_upper:
                found_mes = True
        # Atualiza código corrente se achou um novo
        if found_code:
            current_code = found_code
        # Se achou o mês e tem código corrente, registra
        if found_mes and current_code:
            result[current_code] = row[0].row
    return result

def parse_br_csv(path):
    for enc in ["latin1","utf-8","cp1252"]:
        try: return pd.read_csv(path, encoding=enc, sep=None, engine="python")
        except: pass
    raise ValueError(f"Não foi possível ler {path}")

# ═══════════════════════════════════════════════════════════════════════════════
# PROCESSADORES
# ═══════════════════════════════════════════════════════════════════════════════

def detect_col_map(ws):
    """
    Detecta colunas pelo nome do cabeçalho em qualquer linha das primeiras 10.
    Retorna dict: nome_campo -> número_coluna (1-based).
    """
    CAMPOS = {
        "ValorTabela":    ["ValorTabela", "Valor Tabela", "VALORTABELA"],
        "ValorPraticado": ["ValorPraticado", "Valor Praticado", "VALORPRATICADO"],
        "ValorPedido":    ["ValorPedido", "Valor Pedido", "VALORPEDIDO"],
        "CCResidual":     ["CC Residual", "CCResidual", "CC_RESIDUAL", "Residual"],
        "CMV":            ["CMV"],
        "REM_TRF":        ["REM+TRF", "TRF+REM", "REMTRF", "REM TRF"],
    }
    result = {}
    for row in ws.iter_rows(max_row=10):
        for cell in row:
            if not cell.value: continue
            cs = str(cell.value).strip()
            for campo, aliases in CAMPOS.items():
                if cs in aliases and campo not in result:
                    result[campo] = cell.column
    return result

def processar_cmv(mestre_path, fat_path, cmv_path, pedidos_path, output_path, mes_display):
    mes = MES_PLANILHA[mes_display]
    report = {k:[] for k in ["etapa1_ok","etapa1_nok","etapa2_vd_ok",
                               "etapa2_lojas_ok","etapa2_nok","etapa3_ok","etapa3_nok"]}

    df_fat = parse_br_csv(fat_path)
    # Renomear colunas dinamicamente baseado na posição (layout pode variar)
    COL_NAMES = ["Loja","ReceitaBruta","Descontos","ReceitaLiquida","Trocas","LiquidaMenosTrocas"]
    n_cols = len(df_fat.columns)
    if n_cols >= 6:
        df_fat.columns = COL_NAMES[:n_cols] if n_cols <= len(COL_NAMES) else COL_NAMES + [f"extra_{i}" for i in range(n_cols-len(COL_NAMES))]
        df_fat = df_fat[COL_NAMES]
    elif n_cols >= 4:
        # Fallback: col 0=Loja, col 1=ReceitaBruta, col 2=Descontos, col 3=ReceitaLiquida
        df_fat.columns = COL_NAMES[:n_cols]
        for col in COL_NAMES[n_cols:]:
            df_fat[col] = 0.0
    else:
        raise ValueError(f"GerencialVendas tem apenas {n_cols} colunas — esperado mínimo 4")
    df_fat["CodLoja"] = df_fat["Loja"].apply(store_code)
    for c in ["ReceitaBruta","Descontos","ReceitaLiquida","Trocas","LiquidaMenosTrocas"]:
        df_fat[c] = df_fat[c].apply(br_float)
    df_fat = df_fat.dropna(subset=["CodLoja"])

    df_cmv = parse_br_csv(cmv_path)
    df_cmv["CodLoja"] = df_cmv["Quebra"].apply(store_code)
    df_cmv["CMV Gerencial"] = df_cmv["CMV Gerencial"].apply(br_float)
    df_cmv["Valor TRF REM"] = df_cmv["Valor TRF REM"].apply(br_float)
    df_cmv = df_cmv.dropna(subset=["CodLoja"])

    df_ped = pd.read_excel(pedidos_path)
    cd_map = {}
    for _, r in df_ped[["CodCD","CanalDistribuicao"]].drop_duplicates().iterrows():
        cod = store_code(r["CanalDistribuicao"])
        if cod: cd_map[int(r["CodCD"])] = cod
    df_ped["CodLoja"] = df_ped["CodCD"].map(cd_map)
    agg = df_ped.groupby("CodLoja").agg(
        SomaValorTabela=("ValorTabela","sum"),
        SomaValorPraticado=("ValorPraticado","sum"),
        SomaValorPedido=("ValorPedido","sum"),
        SomaValorTotalSemCCR=("ValorTotalSemCCR","sum"),
    ).reset_index()
    agg["CCResidual"] = agg["SomaValorPedido"] - agg["SomaValorTotalSemCCR"]

    shutil.copy(mestre_path, output_path)
    wb = load_workbook(output_path)
    # Localizar aba VD por possíveis nomes
    VD_NOMES    = ["VD", "VENDA DIRETA", "Venda Direta", "vd", "VENDAS DIRETAS"]
    LOJAS_NOMES = ["LOJAS", "Lojas", "lojas", "LOJA"]
    nome_vd    = next((n for n in VD_NOMES    if n in wb.sheetnames), None)
    nome_lojas = next((n for n in LOJAS_NOMES if n in wb.sheetnames), None)
    if nome_vd is None:
        raise ValueError(f"Aba VD não encontrada. Abas disponíveis: {wb.sheetnames}. Esperado: {VD_NOMES}")
    if nome_lojas is None:
        raise ValueError(f"Aba LOJAS não encontrada. Abas disponíveis: {wb.sheetnames}. Esperado: {LOJAS_NOMES}")
    ws_vd, ws_lojas = wb[nome_vd], wb[nome_lojas]
    map_vd    = build_month_map(ws_vd,    mes)
    map_lojas = build_month_map(ws_lojas, mes)
    # Detectar colunas dinamicamente pelo cabeçalho
    cols_vd    = detect_col_map(ws_vd)
    cols_lojas = detect_col_map(ws_lojas)

    for _, row in df_fat.iterrows():
        code = int(row["CodLoja"])
        if code not in map_lojas: report["etapa1_nok"].append(str(code)); continue
        r = map_lojas[code]
        # Usar colunas detectadas ou fallback para defaults
        _rb = cols_lojas.get("ReceitaBruta",    COL_FAT["ReceitaBruta"])
        _dc = cols_lojas.get("Descontos",       COL_FAT["Descontos"])
        _rl = cols_lojas.get("ReceitaLiquida",  COL_FAT["ReceitaLiquida"])
        _tr = cols_lojas.get("Trocas",          COL_FAT["Trocas"])
        _lt = cols_lojas.get("LiquidaMenosTrocas", COL_FAT["LiquidaMenosTrocas"])
        ws_lojas.cell(row=r, column=_rb).value = row["ReceitaBruta"]
        ws_lojas.cell(row=r, column=_dc).value = row["Descontos"]
        ws_lojas.cell(row=r, column=_rl).value = row["ReceitaLiquida"]
        ws_lojas.cell(row=r, column=_tr).value = row["Trocas"]
        ws_lojas.cell(row=r, column=_lt).value = row["LiquidaMenosTrocas"]
        report["etapa1_ok"].append(str(code))

    for _, row in df_cmv.iterrows():
        code = int(row["CodLoja"]); cmv_val = row["CMV Gerencial"]; trf_val = row["Valor TRF REM"]
        found = False
        if code in map_vd:
            r = map_vd[code]
            cmv_c = cols_vd.get("CMV",    COL_CMV_VD["CMV"])
            trf_c = cols_vd.get("REM_TRF",COL_CMV_VD["TRF"])
            ws_vd.cell(row=r, column=cmv_c).value = cmv_val
            ws_vd.cell(row=r, column=trf_c).value = trf_val
            report["etapa2_vd_ok"].append(str(code)); found = True
        if code in map_lojas:
            r = map_lojas[code]
            cmv_c = cols_lojas.get("CMV",    COL_CMV_LOJAS["CMV"])
            trf_c = cols_lojas.get("REM_TRF",COL_CMV_LOJAS["TRF"])
            ws_lojas.cell(row=r, column=cmv_c).value = cmv_val
            ws_lojas.cell(row=r, column=trf_c).value = trf_val
            report["etapa2_lojas_ok"].append(str(code)); found = True
        if not found: report["etapa2_nok"].append(str(code))

    for _, row in agg.iterrows():
        code = int(row["CodLoja"])
        if code not in map_vd: report["etapa3_nok"].append(str(code)); continue
        r = map_vd[code]
        if cols_vd.get("ValorTabela"):    ws_vd.cell(row=r, column=cols_vd["ValorTabela"]).value    = round(row["SomaValorTabela"],2)
        if cols_vd.get("ValorPraticado"): ws_vd.cell(row=r, column=cols_vd["ValorPraticado"]).value = round(row["SomaValorPraticado"],2)
        if cols_vd.get("ValorPedido"):    ws_vd.cell(row=r, column=cols_vd["ValorPedido"]).value    = round(row["SomaValorPedido"],2)
        if cols_vd.get("CCResidual"):     ws_vd.cell(row=r, column=cols_vd["CCResidual"]).value     = round(row["CCResidual"],2)
        report["etapa3_ok"].append(str(code))

    wb.save(output_path)
    return report


def processar_estoque(mestre_path, posicao_path, nf_compra_path, cmv_path, output_path, mes_display):
    aba = MES_ABA_ESTOQUE[mes_display]
    report = {"estoque_ok":[],"estoque_nok":[],"compras_ok":[],"compras_nok":[],"cmv_ok":[],"cmv_nok":[]}

    df_pos = parse_br_csv(posicao_path)
    df_pos["SaldoAtual"] = df_pos["Saldo Atual"].astype(str).str.replace(",",".").astype(float)
    df_pos["CustoF"]     = df_pos["Custo"].astype(str).str.replace(",",".").astype(float)
    df_pos["ValorEst"]   = df_pos["SaldoAtual"] * df_pos["CustoF"]
    est_por_loja = df_pos.groupby("Quebra")["ValorEst"].sum()

    df_nf = parse_br_csv(nf_compra_path)
    df_nf["VT"] = df_nf["Valor Total"].astype(str).str.replace(".","").str.replace(",",".").astype(float)
    compras_por_loja = df_nf.groupby("Código da Loja")["VT"].sum()

    df_cmv = parse_br_csv(cmv_path)
    df_cmv["CodLoja"] = df_cmv["Quebra"].apply(store_code)
    df_cmv["CMV Gerencial"] = df_cmv["CMV Gerencial"].apply(br_float)
    df_cmv = df_cmv.dropna(subset=["CodLoja"])
    cmv_por_loja = df_cmv.set_index("CodLoja")["CMV Gerencial"]

    shutil.copy(mestre_path, output_path)
    wb = load_workbook(output_path)
    ws = wb[aba]

    for row in ws.iter_rows(min_row=9, max_col=1):
        cell = row[0]
        if cell.value is None: continue
        code = store_code(str(cell.value))
        if not code: continue
        r = cell.row
        if code in est_por_loja.index:
            ws.cell(row=r,column=2).value = round(float(est_por_loja[code]),2)
            report["estoque_ok"].append(str(code))
        else: report["estoque_nok"].append(str(code))
        if code in compras_por_loja.index:
            ws.cell(row=r,column=3).value = round(float(compras_por_loja[code]),2)
            report["compras_ok"].append(str(code))
        else: report["compras_nok"].append(str(code))
        if code in cmv_por_loja.index:
            ws.cell(row=r,column=4).value = round(float(cmv_por_loja[code]),2)
            report["cmv_ok"].append(str(code))
        else: report["cmv_nok"].append(str(code))

    wb.save(output_path)
    return report



def extrair_linhas_fc(fc_path):
    """Lê col D do FC (Real 2022) e retorna {descricao: excel_row} para linhas não vazias."""
    _xl_tmp = pd.ExcelFile(fc_path)
    _aba_real = next((s for s in _xl_tmp.sheet_names if "real" in s.lower()), "Real 2022")
    df = pd.read_excel(fc_path, sheet_name=_aba_real, header=None)
    linhas = {}
    for i in range(len(df)):
        v = df.iloc[i, 3]  # col D
        if pd.notna(v) and str(v).strip():
            linhas[str(v).strip()] = i + 1  # excel row (1-based)
    return linhas

def extrair_grupos_plano(fc_path):
    """Lê col B do Plano de Contas F360 e retorna grupos únicos (excluindo ignorados)."""
    df = pd.read_excel(fc_path, sheet_name="Plano de Contas F360", header=0)
    grupos = df["Fornecedor/Grupo"].dropna().astype(str).str.strip().unique().tolist()
    ignorar = {"nan","não","nao","não entra","nao entra","nãoentra","naoentra","não entra"}
    return [g for g in grupos if g and g.lower() not in ignorar]

def processar_fc(mestre_path, posicao_path, nf_compra_path, cmv_path,
                 fat_path, pedidos_path, rateio_path, cartoes_paths,
                 output_path, mes_display, mapeamentos_extra=None,
                 grupo_linha_map=None):
    ano    = "26"
    prefixo    = MES_FC_PREFIX[mes_display]
    header_mes = f"{prefixo}'{ano}"

    report = {
        "mes_col":None,
        "f1_boletos":0.0,"f2_cartoes":0.0,"f3_estoque":0.0,
        "f4_fornecedores":{},"f4_alertas":[],
        "f5_receita":0.0,"f7_cmv":0.0,"f6_compras":0.0,
        "plano_contas_atualizado":[],
    }

    # F3: total estoque
    df_pos = parse_br_csv(posicao_path)
    df_pos["SaldoAtual"] = df_pos["Saldo Atual"].astype(str).str.replace(",",".").astype(float)
    df_pos["CustoF"]     = df_pos["Custo"].astype(str).str.replace(",",".").astype(float)
    report["f3_estoque"] = round((df_pos["SaldoAtual"] * df_pos["CustoF"]).sum(), 2)

    # F6: total NF Compras
    df_nf = parse_br_csv(nf_compra_path)
    df_nf["VT"] = df_nf["Valor Total"].astype(str).str.replace(".","").str.replace(",",".").astype(float)
    report["f6_compras"] = round(df_nf["VT"].sum(), 2)

    # F5: Receita Líquida = GerencialVendas col F + ConsultaPedidos ValorPedido
    df_fat = parse_br_csv(fat_path)
    _col_names_fc = ["Loja","ReceitaBruta","Descontos","ReceitaLiquida","Trocas","LiquidaMenosTrocas"]
    _n = len(df_fat.columns)
    df_fat.columns = _col_names_fc[:_n] if _n <= len(_col_names_fc) else _col_names_fc + [f"extra_{i}" for i in range(_n-len(_col_names_fc))]
    for _c in _col_names_fc:
        if _c not in df_fat.columns: df_fat[_c] = 0.0
    total_fat = sum(br_float(v) or 0 for v in df_fat["LiquidaMenosTrocas"])
    df_ped = pd.read_excel(pedidos_path)
    report["f5_receita"] = round(total_fat + float(df_ped["ValorPedido"].sum()), 2)

    # F7: CMV do mês = soma total CMV Gerencial
    df_cmv_fc = parse_br_csv(cmv_path)
    df_cmv_fc["CMV Gerencial"] = df_cmv_fc["CMV Gerencial"].apply(br_float)
    report["f7_cmv"] = round(df_cmv_fc["CMV Gerencial"].dropna().sum(), 2)

    # F4: Rateio → De-Para
    df_rat = pd.read_excel(rateio_path, header=9)
    df_rat["Valor Bruto"] = pd.to_numeric(df_rat["Valor Bruto"], errors="coerce")
    # Excluir Tipo de Documento = 'Previsão' e Status = 'Baixado'
    if "Tipo de Documento" in df_rat.columns:
        df_rat = df_rat[df_rat["Tipo de Documento"].astype(str).str.strip() != "Previsão"]
    if "Status" in df_rat.columns:
        df_rat = df_rat[df_rat["Status"].astype(str).str.strip() != "Baixado"]
    df_depara = pd.read_excel(mestre_path, sheet_name="Plano de Contas F360",
                               engine="openpyxl", header=0)
    depara_map = {
        str(k).strip(): str(v).strip()
        for k, v in zip(df_depara["Categoria"], df_depara["Fornecedor/Grupo"])
        if pd.notna(k) and pd.notna(v)
    }
    # Incorporar mapeamentos extras do usuário
    if mapeamentos_extra:
        for conta_extra, grupo_extra in mapeamentos_extra.items():
            if grupo_extra and grupo_extra != "— Ignorar —":
                depara_map[conta_extra] = grupo_extra
    soma_por_grupo = {}
    alertas = []
    # grupo_linha_map: {nome_grupo_plano: linha_excel_fc} — definido pelo usuário
    _glmap = grupo_linha_map or {}
    for _, row in df_rat.iterrows():
        conta = str(row["Plano de Contas"]).strip() if pd.notna(row["Plano de Contas"]) else ""
        valor = row["Valor Bruto"] if pd.notna(row["Valor Bruto"]) else 0.0
        if not conta: continue
        grupo = depara_map.get(conta)
        if grupo is None:
            if conta not in alertas: alertas.append(conta)
            continue
        if grupo.upper() in ("NÃO", "NAO", "NÃO ENTRA", "NAO ENTRA", "— IGNORAR —"):
            continue
        linha_excel = _glmap.get(grupo) or _glmap.get(grupo.strip())
        if linha_excel is None:
            if grupo not in alertas: alertas.append(f"Grupo sem linha FC: {grupo}")
            continue
        soma_por_grupo[linha_excel] = soma_por_grupo.get(linha_excel, 0.0) + valor
    report["f4_fornecedores"] = {k:round(v,2) for k,v in soma_por_grupo.items()}
    report["f4_alertas"] = alertas

    # F1/F2: Parcelas de Cartões
    frames = [pd.read_excel(cp, header=0) for cp in cartoes_paths]
    df_cart = pd.concat(frames, ignore_index=True)
    # Excluir onde col Z (index 25) está preenchida
    obs_col = df_cart.iloc[:, 25].astype(str).str.strip()
    df_cart = df_cart[(obs_col == "") | (obs_col == "nan") | (obs_col.isna())]
    df_cart["VB"] = pd.to_numeric(df_cart.iloc[:, 11], errors="coerce").fillna(0)
    modalidade = df_cart.iloc[:, 7].astype(str).str.strip()
    report["f1_boletos"] = round(float(df_cart[modalidade == "Boleto Garantido"]["VB"].sum()), 2)
    report["f2_cartoes"] = round(float(df_cart[modalidade != "Boleto Garantido"]["VB"].sum()), 2)

    # Escrever no workbook - sempre salvar como xlsx (openpyxl não suporta VBA)
    output_path = output_path.replace(".xlsm", ".xlsx")
    shutil.copy(mestre_path, output_path)
    wb = load_workbook(output_path)
    ws = wb["Real 2022"]

    mes_col = None
    for cell in ws[7]:
        if cell.value and str(cell.value).strip().upper() == header_mes.upper():
            mes_col = cell.column; break

    if mes_col is None:
        report["mes_col"] = f"❌ Coluna '{header_mes}' não encontrada na linha 7!"
        wb.save(output_path)
        return report

    report["mes_col"] = f"Coluna {mes_col} ({header_mes})"
    ws.cell(row=11, column=mes_col).value = report["f1_boletos"]
    ws.cell(row=12, column=mes_col).value = report["f2_cartoes"]
    ws.cell(row=13, column=mes_col).value = report["f3_estoque"]
    for linha, valor in report["f4_fornecedores"].items():
        ws.cell(row=linha, column=mes_col).value = valor
    ws.cell(row=37, column=mes_col).value = report["f5_receita"]
    ws.cell(row=38, column=mes_col).value = report["f7_cmv"]
    ws.cell(row=39, column=mes_col).value = report["f6_compras"]

    # Atualizar aba Plano de Contas F360 com novos mapeamentos do usuário
    if mapeamentos_extra:
        ws_pc = wb["Plano de Contas F360"]
        # Encontrar próxima linha vazia na col A (após o cabeçalho)
        ultima_linha = 1
        for row_pc in ws_pc.iter_rows(min_col=1, max_col=1):
            if row_pc[0].value is not None:
                ultima_linha = row_pc[0].row
        contas_ja_existentes = set()
        for row_pc in ws_pc.iter_rows(min_col=1, max_col=1, min_row=2):
            if row_pc[0].value:
                contas_ja_existentes.add(str(row_pc[0].value).strip())
        novas_contas = []
        for conta_nova, grupo_novo in mapeamentos_extra.items():
            if grupo_novo and grupo_novo != "— Ignorar —" and conta_nova not in contas_ja_existentes:
                proxima = ultima_linha + 1
                ws_pc.cell(row=proxima, column=1).value = conta_nova
                ws_pc.cell(row=proxima, column=2).value = grupo_novo
                ultima_linha = proxima
                novas_contas.append(f"{conta_nova} → {grupo_novo}")
        report["plano_contas_atualizado"] = novas_contas

    wb.save(output_path)
    return report


# ═══════════════════════════════════════════════════════════════════════════════
# STREAMLIT
# ═══════════════════════════════════════════════════════════════════════════════
def run_streamlit():
    import streamlit as st
    import tempfile

    st.set_page_config(page_title="VENA | Atualização Mensal", page_icon="🧬", layout="wide")
    st.markdown("""
    <style>
        .stApp{background-color:#f4f6fb}
        .vena-header{background:linear-gradient(90deg,#4A90D9,#5B9FE8);padding:18px 36px;
            border-radius:0 0 16px 16px;display:flex;align-items:center;gap:16px;
            margin-bottom:28px;box-shadow:0 4px 16px rgba(74,144,217,.18)}
        .vena-logo{background:white;border-radius:10px;padding:6px 16px;
            font-size:20px;font-weight:900;color:#4A90D9;letter-spacing:2px}
        .vena-header-title{color:white;font-size:18px;font-weight:700}
        .vena-header-sub{color:rgba(255,255,255,.85);font-size:12px;margin-top:2px}
        .section-title{color:#2d5fa6;font-size:13px;font-weight:700;margin:20px 0 10px;
            padding-bottom:6px;border-bottom:2px solid #e0e8f5;text-transform:uppercase;letter-spacing:.5px}
        .upload-card{background:white;border-radius:14px;padding:16px 20px;
            box-shadow:0 2px 10px rgba(74,144,217,.10);margin-bottom:14px;border-left:5px solid #4A90D9}
        .upload-card h4{color:#2d5fa6;font-size:13px;font-weight:700;margin-bottom:8px}
        .mes-card{background:white;border-radius:14px;padding:16px 20px;
            box-shadow:0 2px 10px rgba(123,94,167,.12);margin-bottom:20px;border-left:5px solid #7B5EA7}
        .mes-card h4{color:#7B5EA7;font-size:13px;font-weight:700;margin-bottom:8px}
        .stButton>button{background:linear-gradient(135deg,#7B5EA7,#9B6FD4)!important;
            color:white!important;border:none!important;border-radius:10px!important;
            padding:14px!important;font-size:16px!important;font-weight:700!important;
            box-shadow:0 4px 16px rgba(123,94,167,.30)!important;width:100%}
        [data-testid="metric-container"]{background:white;border-radius:12px;padding:16px 20px;
            box-shadow:0 2px 10px rgba(74,144,217,.10);border-top:4px solid #7B5EA7}
        [data-testid="stMetricValue"]{color:#2d5fa6;font-weight:800}
        .stDownloadButton>button{background:linear-gradient(135deg,#4A90D9,#5B9FE8)!important;
            color:white!important;border-radius:10px!important;font-weight:700!important;
            border:none!important;width:100%;padding:12px!important;font-size:15px!important}
        .alert-box{background:#fff3cd;border:1px solid #ffc107;border-radius:10px;
            padding:12px 16px;color:#856404;font-size:13px;margin-top:8px}
        hr{border-color:#e0e8f5}
    </style>
    """, unsafe_allow_html=True)

    st.markdown("""
    <div class="vena-header">
        <div class="vena-logo">VENA</div>
        <div>
            <div class="vena-header-title">Atualização Mensal 2026</div>
            <div class="vena-header-sub">CMV · Estoque · FC Consolidado</div>
        </div>
    </div>
    """, unsafe_allow_html=True)

    # Mês
    st.markdown('<div class="section-title">📅 Mês de Referência</div>', unsafe_allow_html=True)
    st.markdown('<div class="mes-card"><h4>📅 Selecione o mês que será atualizado</h4></div>',
                unsafe_allow_html=True)
    mes_display = st.selectbox("Mês", MESES_DISPLAY, index=2, label_visibility="collapsed")

    modo_atualizar = True

    st.markdown('<div class="section-title">📋 Planilhas para Atualizar</div>', unsafe_allow_html=True)
    c1,c2,c3 = st.columns(3)
    atualizar_cmv     = c1.checkbox("📊 CMV Mensal",             value=True)
    atualizar_estoque = c2.checkbox("📦 Informativo de Estoque", value=True)
    atualizar_fc      = c3.checkbox("💰 FC Consolidado",         value=True)
    atualizar_espelho = st.checkbox("📋 Espelho DFC",             value=False)
    atualizar_dfc = False
    validar_dfc   = False

    st.markdown('<div class="section-title">📝 Lançamentos Manuais (ValoresDaDRE)</div>', unsafe_allow_html=True)
    lm1, lm2, lm3 = st.columns(3)
    lanc_vendas = lm1.checkbox("🛍️ Vendas de Mercadorias", value=False)
    st.markdown("---")

    # Planilhas Mestres (aparecem só se checkbox ativo)
    cmv_mestre = estoque_mestre = fc_mestre = dfc_mestre = None
    dre_vendas_mestre = None
    # espelho_dfc não tem mestre — gera novo arquivo
    if atualizar_cmv or atualizar_estoque or atualizar_fc or lanc_vendas:
        st.markdown('<div class="section-title">📁 Planilhas Mestres</div>', unsafe_allow_html=True)
        mc1,mc2,mc3 = st.columns(3)
        with mc1:
            if atualizar_cmv or atualizar_espelho or lanc_vendas:
                st.markdown('<div class="upload-card"><h4>📊 CMV Mensal (.xlsx)</h4></div>',
                            unsafe_allow_html=True)
                cmv_mestre = st.file_uploader("CMV Mestre", type=["xlsx"],
                                              label_visibility="collapsed", key="cmv_m")
            if lanc_vendas:
                st.markdown('<div class="upload-card"><h4>🛍️ DRE — Vendas de Mercadorias (.xlsx)</h4></div>',
                            unsafe_allow_html=True)
                dre_vendas_mestre = st.file_uploader("DRE Vendas", type=["xlsx"],
                                                      label_visibility="collapsed", key="dre_vend")
        with mc2:
            if atualizar_estoque or atualizar_dfc or atualizar_espelho:
                st.markdown('<div class="upload-card"><h4>📦 Informativo de Estoque (.xlsx)</h4></div>',
                            unsafe_allow_html=True)
                estoque_mestre = st.file_uploader("Estoque Mestre", type=["xlsx"],
                                                   label_visibility="collapsed", key="est_m")
        with mc3:

            if atualizar_dfc:
                st.markdown('<div class="upload-card"><h4>📈 DFC (.xlsb)</h4></div>',
                            unsafe_allow_html=True)
                dfc_mestre = st.file_uploader("DFC Mestre", type=["xlsb"],
                                               label_visibility="collapsed", key="dfc_m")
            if atualizar_fc or atualizar_dfc or atualizar_espelho:
                st.markdown('<div class="upload-card"><h4>💰 FC Consolidado (.xlsm)</h4></div>',
                            unsafe_allow_html=True)
                fc_mestre = st.file_uploader("FC Mestre", type=["xlsm","xlsx"],
                                              label_visibility="collapsed", key="fc_m")

    # Bases de dados (aparecem apenas se necessárias)
    st.markdown('<div class="section-title">📂 Bases de Dados</div>', unsafe_allow_html=True)
    bc1, bc2 = st.columns(2)
    posicao_file = nf_file = cmv_csv_file = fat_file = None
    pedidos_file = rateio_file = cartoes_files = rateio_ant_file = cartoes_ant_files = None

    with bc1:
        if atualizar_estoque or atualizar_fc:
            st.markdown('<div class="upload-card"><h4>📦 Posição de Estoque (CSV)</h4></div>',
                        unsafe_allow_html=True)
            posicao_file = st.file_uploader("Posição", type=["csv"],
                                             label_visibility="collapsed", key="pos")
        if atualizar_cmv or atualizar_estoque or lanc_cmv:
            st.markdown('<div class="upload-card"><h4>📉 CMV Gerencial (CSV)</h4></div>',
                        unsafe_allow_html=True)
            cmv_csv_file = st.file_uploader("CMV CSV", type=["csv"],
                                             label_visibility="collapsed", key="cmv_csv")
        if atualizar_cmv or atualizar_fc or lanc_vendas:
            st.markdown('<div class="upload-card"><h4>🛒 Faturamento VD (.xlsx)</h4></div>',
                        unsafe_allow_html=True)
            pedidos_file = st.file_uploader("Faturamento VD", type=["xlsx"],
                                             label_visibility="collapsed", key="ped")
        if atualizar_espelho:
            st.markdown('<div class="upload-card"><h4>📑 Contas a Pagar — Mês Anterior (.xlsx)</h4></div>',
                        unsafe_allow_html=True)
            rateio_ant_file = st.file_uploader("Contas a Pagar Anterior", type=["xlsx"],
                                               label_visibility="collapsed", key="rat_ant")
        if atualizar_fc:
            st.markdown('<div class="upload-card"><h4>📑 Contas a Pagar (.xlsx)</h4></div>',
                        unsafe_allow_html=True)
            rateio_file = st.file_uploader("Contas a Pagar", type=["xlsx"],
                                            label_visibility="collapsed", key="rat")

    with bc2:
        if atualizar_estoque or atualizar_fc:
            st.markdown('<div class="upload-card"><h4>🧾 NF Compras Sintético (CSV)</h4></div>',
                        unsafe_allow_html=True)
            nf_file = st.file_uploader("NF Compras", type=["csv"],
                                        label_visibility="collapsed", key="nf")
        if atualizar_cmv or atualizar_fc or lanc_vendas:
            st.markdown('<div class="upload-card"><h4>💰 Faturamento Loja — GerencialVendas (CSV)</h4></div>',
                        unsafe_allow_html=True)
            fat_file = st.file_uploader("Faturamento", type=["csv"],
                                         label_visibility="collapsed", key="fat")
        if atualizar_fc:
            st.markdown('<div class="upload-card"><h4>💳 Contas a Receber (.xlsx) — múltiplos permitidos</h4></div>',
                        unsafe_allow_html=True)
            cartoes_files = st.file_uploader("Contas a Receber", type=["xlsx"], accept_multiple_files=True,
                                              label_visibility="collapsed", key="cart")
        if atualizar_espelho:
            st.markdown('<div class="upload-card"><h4>💳 Contas a Receber — Mês Anterior (.xlsx) — múltiplos</h4></div>',
                        unsafe_allow_html=True)
            cartoes_ant_files = st.file_uploader("Contas a Receber Anterior", type=["xlsx"],
                                                  accept_multiple_files=True,
                                                  label_visibility="collapsed", key="cart_ant")

    # Validação
    erros = []
    if atualizar_cmv     and not cmv_mestre:      erros.append("CMV Mestre")
    if atualizar_estoque and not estoque_mestre:   erros.append("Estoque Mestre")
    if (atualizar_fc or atualizar_dfc) and not fc_mestre: erros.append("FC Consolidado")
    if (atualizar_estoque or atualizar_fc) and not posicao_file:  erros.append("Posição de Estoque")
    if (atualizar_estoque or atualizar_fc) and not nf_file:       erros.append("NF Compras")
    if (atualizar_cmv or atualizar_estoque or atualizar_fc) and not cmv_csv_file: erros.append("CMV Gerencial CSV")
    if (atualizar_cmv or atualizar_fc) and not fat_file:          erros.append("Faturamento")
    if (atualizar_cmv or atualizar_fc) and not pedidos_file:      erros.append("Faturamento VD")
    if atualizar_fc and not rateio_file:   erros.append("Contas a Pagar")
    if atualizar_fc and not cartoes_files: erros.append("Contas a Receber")

    if lanc_vendas and not dre_vendas_mestre: erros.append("DRE Vendas de Mercadorias")
    if lanc_cmv    and not dre_cmv_mestre:    erros.append("DRE CMV")
    if lanc_cmv    and not cmv_csv_file:       erros.append("CMV Gerencial CSV (para CMV)")
    if lanc_vendas and not cmv_mestre:        erros.append("CMV Mestre (para identificar canal VD/Loja)")
    if lanc_vendas and not fat_file:          erros.append("Faturamento Loja (para Vendas de Mercadorias)")
    if lanc_vendas and not pedidos_file:      erros.append("Faturamento VD (para Vendas de Mercadorias)")
    if erros: st.info(f"📎 Pendentes: **{', '.join(erros)}**")
    else:      st.success(f"✅ Todos os arquivos prontos para processar **{mes_display}**.")
    st.markdown("---")

    # ── Pré-verificação de contas não mapeadas (FC) ─────────────────────────
    if atualizar_fc and not erros and rateio_file and fc_mestre:
        import io as _io
        rateio_bytes = rateio_file.read(); rateio_file.seek(0)
        fc_mestre.seek(0)
        fc_bytes = fc_mestre.read()
        # Salvar bytes no session_state para uso posterior mesmo após cursor consumido
        st.session_state["_fc_bytes"] = fc_bytes
        fc_mestre.seek(0)
        df_rat_pre   = pd.read_excel(_io.BytesIO(rateio_bytes), header=9)
        # Aplicar os mesmos filtros da função principal
        if "Tipo de Documento" in df_rat_pre.columns:
            df_rat_pre = df_rat_pre[df_rat_pre["Tipo de Documento"].astype(str).str.strip() != "Previsão"]
        if "Status" in df_rat_pre.columns:
            df_rat_pre = df_rat_pre[df_rat_pre["Status"].astype(str).str.strip() != "Baixado"]
        # Verificar se aba Plano de Contas F360 existe
        _xl_pre = pd.ExcelFile(_io.BytesIO(fc_bytes))
        if "Plano de Contas F360" not in _xl_pre.sheet_names:
            st.warning(f"⚠️ O FC carregado não tem aba 'Plano de Contas F360'. Abas disponíveis: {_xl_pre.sheet_names}")
            contas_sem_mapa = []
            dep_map_pre = {}
        else:
            df_dep_pre   = pd.read_excel(_io.BytesIO(fc_bytes),
                                         sheet_name="Plano de Contas F360", header=0)
            dep_map_pre  = dict(zip(df_dep_pre["Categoria"].astype(str).str.strip(),
                                    df_dep_pre["Fornecedor/Grupo"].astype(str).str.strip()))
        if dep_map_pre:
            contas_sem_mapa = []
            for conta in df_rat_pre["Plano de Contas"].dropna().astype(str).str.strip().unique():
                if conta and conta not in dep_map_pre:
                    contas_sem_mapa.append(conta)

        if contas_sem_mapa:
            st.markdown("---")
            st.markdown('<div class="section-title">⚠️ Contas Sem Mapeamento — Classificar antes de processar</div>',
                        unsafe_allow_html=True)
            st.warning("As contas abaixo não estão no Plano de Contas F360. Selecione o agrupamento de cada uma antes de continuar.")
            # Usar grupos reais do FC carregado (col B do Plano de Contas F360)
            _grupos_reais = sorted(dep_map_pre.values())
            _grupos_unicos = list(dict.fromkeys(
                g for g in _grupos_reais
                if g and g.upper() not in ("NÃO ENTRA", "NAO ENTRA", "NÃO", "NAO", "NAN")
            ))
            GRUPOS_OPCOES = ["— Ignorar —"] + _grupos_unicos
            mapeamentos_usuario = {}
            for i, conta in enumerate(contas_sem_mapa):
                col_conta, col_sel = st.columns([3,2])
                col_conta.markdown(f"**{conta}**")
                grupo_escolhido = col_sel.selectbox(
                    "Agrupamento", GRUPOS_OPCOES,
                    key=f"grupo_{i}", label_visibility="collapsed"
                )
                mapeamentos_usuario[conta] = grupo_escolhido
            st.session_state["mapeamentos_usuario"] = mapeamentos_usuario
            st.markdown("---")
        else:
            st.session_state["mapeamentos_usuario"] = {}

    mapeamentos_usuario = st.session_state.get("mapeamentos_usuario", {})

    # ── Mapeamento Grupo → Linha FC ───────────────────────────────────────────
    # Salvar grupos e linhas no session_state quando FC é carregado
    # para que fiquem disponíveis mesmo após o botão ser pressionado
    if atualizar_fc and fc_mestre and not erros:
        import io as _io2
        fc_mestre.seek(0)
        fc_bytes2 = fc_mestre.read()
        if fc_bytes2:
            st.session_state["_fc_bytes"] = fc_bytes2
        fc_mestre.seek(0)
        try:
            _grupos = extrair_grupos_plano(_io2.BytesIO(fc_bytes2))
            _linhas = extrair_linhas_fc(_io2.BytesIO(fc_bytes2))
            # Persistir no session_state para sobreviver ao clique do botão
            st.session_state["_fc_grupos"] = _grupos
            st.session_state["_fc_linhas"] = _linhas
        except Exception as _e:
            st.warning(f"Não foi possível ler grupos do FC: {_e}")

    # Ler grupos/linhas do session_state (persiste após clique do botão)
    _grupos = st.session_state.get("_fc_grupos", [])
    _linhas = st.session_state.get("_fc_linhas", {})
    # Normalizar chaves com strip para evitar erros de espaço
    _linhas  = {k.strip(): v for k, v in _linhas.items()}
    _grupos  = [g.strip() for g in _grupos]
    _opcoes  = ["— Ignorar —"] + list(_linhas.keys())

    if atualizar_fc and _grupos:
        st.markdown("---")
        st.markdown('<div class="section-title">🔗 Mapeamento: Fornecedor/Grupo → Linha do FC</div>',
                    unsafe_allow_html=True)
        st.caption("Indique a qual linha do FC (aba Real 2022) cada grupo do Plano de Contas corresponde.")

        cols_map = st.columns(2)
        for idx, grupo in enumerate(_grupos):
            col = cols_map[idx % 2]
            # Default: match automático por nome (após strip), senão "— Ignorar —"
            grupo_s = grupo.strip()
            default_name = grupo_s if grupo_s in _opcoes else "— Ignorar —"
            # Se já foi selecionado antes, manter a escolha salva
            salvo = st.session_state.get(f"glmap_{idx}")
            if salvo and salvo in _opcoes:
                default_name = salvo
            col.selectbox(
                f"**{grupo}**", _opcoes,
                index=_opcoes.index(default_name),
                key=f"glmap_{idx}"
            )

    # Reconstruir mapa a partir dos valores ATUAIS dos selectboxes
    grupo_linha_usuario = {}
    for idx, grupo in enumerate(_grupos):
        escolha = st.session_state.get(f"glmap_{idx}", "— Ignorar —")
        if escolha and escolha != "— Ignorar —" and escolha in _linhas:
            grupo_linha_usuario[grupo] = _linhas[escolha]


    if st.button(f"🚀  Processar {mes_display}", disabled=bool(erros)):
        with tempfile.TemporaryDirectory() as tmpdir:
            def save_tmp(f, name):
                p = os.path.join(tmpdir, name)
                with open(p,"wb") as o: o.write(f.read())
                return p

            paths = {}
            # Save FC mestre to disk immediately — cursor may be exhausted from pre-checks
            # Use session_state bytes if available as fallback
            if fc_mestre:
                _fc_ext0 = "xlsm" if fc_mestre.name.endswith("xlsm") else "xlsx"
                _fc_base = os.path.join(tmpdir, f"fc_base.{_fc_ext0}")
                # Try reading fresh; if empty use stored bytes
                fc_mestre.seek(0)
                _fc_raw = fc_mestre.read()
                if not _fc_raw and st.session_state.get("_fc_bytes"):
                    _fc_raw = st.session_state["_fc_bytes"]
                if _fc_raw:
                    with open(_fc_base, "wb") as _f0:
                        _f0.write(_fc_raw)
                    paths["fc_saved"] = _fc_base
            if posicao_file:  paths["pos"]     = save_tmp(posicao_file,  "pos.csv")
            if nf_file:       paths["nf"]      = save_tmp(nf_file,       "nf.csv")
            if cmv_csv_file:  paths["cmv_csv"] = save_tmp(cmv_csv_file,  "cmv.csv")
            if fat_file:      paths["fat"]     = save_tmp(fat_file,      "fat.csv")
            if pedidos_file:  paths["ped"]     = save_tmp(pedidos_file,  "ped.xlsx")
            if rateio_file:   paths["rat"]     = save_tmp(rateio_file,   "rat.xlsx")
            if cartoes_files:     paths["cart"]     = [save_tmp(f,f"cart_{i}.xlsx") for i,f in enumerate(cartoes_files)]
            if rateio_ant_file:   paths["rat_ant"]   = save_tmp(rateio_ant_file, "rat_ant.xlsx")
            if cartoes_ant_files: paths["cart_ant"]  = [save_tmp(f,f"cart_ant_{i}.xlsx") for i,f in enumerate(cartoes_ant_files)]

            # Salvar no session_state — persiste após clique de download
            st.session_state["resultados"]     = {}
            st.session_state["mes_processado"] = mes_display

            if atualizar_cmv:
                with st.spinner("⚙️ CMV..."):
                    try:
                        out = os.path.join(tmpdir,"cmv_out.xlsx")
                        mp  = save_tmp(cmv_mestre,"cmv_m.xlsx")
                        rep = processar_cmv(mp,paths["fat"],paths["cmv_csv"],paths["ped"],out,mes_display)
                        with open(out,"rb") as f:
                            st.session_state["resultados"]["cmv"] = (f.read(), rep)
                    except Exception as e:
                        st.error(f"❌ Erro no CMV: {e}")
                        st.stop()

            if atualizar_estoque:
                with st.spinner("⚙️ Estoque..."):
                    try:
                        out = os.path.join(tmpdir,"est_out.xlsx")
                        mp  = save_tmp(estoque_mestre,"est_m.xlsx")
                        rep = processar_estoque(mp,paths["pos"],paths["nf"],paths["cmv_csv"],out,mes_display)
                        with open(out,"rb") as f:
                            st.session_state["resultados"]["estoque"] = (f.read(), rep)
                    except Exception as e:
                        st.error(f"❌ Erro no Estoque: {e}")
                        st.stop()

            if atualizar_fc:
                with st.spinner("⚙️ FC Consolidado..."):
                    try:
                        ext = "xlsx"
                        out = os.path.join(tmpdir, "fc_out.xlsx")
                        # Usar arquivo já salvo no início do tmpdir
                        _fc_v2 = paths.get("fc_saved")
                        if not _fc_v2 or not os.path.exists(_fc_v2):
                            raise ValueError("FC Consolidado não pôde ser lido — tente recarregar o arquivo")
                        mp = _fc_v2
                        # Resolver paths — usar já salvos ou reler uploads
                        def _resolve(key, file, name):
                            return paths.get(key) or (save_tmp(file, name) if file else None)
                        _pos  = _resolve("pos",     posicao_file,  "pos2.csv")
                        _nf   = _resolve("nf",      nf_file,       "nf2.csv")
                        _csv  = _resolve("cmv_csv", cmv_csv_file,  "cmv2.csv")
                        _fat  = _resolve("fat",     fat_file,       "fat2.csv")
                        _ped  = _resolve("ped",     pedidos_file,  "ped2.xlsx")
                        _rat  = _resolve("rat",     rateio_file,   "rat2.xlsx")
                        _cart = paths.get("cart") or ([save_tmp(f,f"cart2_{i}.xlsx") for i,f in enumerate(cartoes_files)] if cartoes_files else [])
                        # Verificar quais obrigatórios estão faltando
                        _faltando = [n for n,v in [("Posição de Estoque",_pos),("NF Compras",_nf),
                            ("CMV Gerencial",_csv),("Faturamento Loja",_fat),
                            ("Faturamento VD",_ped),("Contas a Pagar",_rat)] if not v]
                        if _faltando:
                            raise ValueError(f"Arquivos obrigatórios para o FC não carregados: {', '.join(_faltando)}")
                        # Ler mapeamento mais recente do session_state
                        # Reconstruir mapa diretamente das keys glmap_N do session_state
                        # (mais confiável que grupo_linha_usuario que pode estar desatualizado)
                        _grupos_ss  = st.session_state.get("_fc_grupos", [])
                        _linhas_ss  = st.session_state.get("_fc_linhas", {})
                        _linhas_ss  = {k.strip(): v for k, v in _linhas_ss.items()}
                        _glmap_atual = {}
                        for _idx, _grupo in enumerate([g.strip() for g in _grupos_ss]):
                            _escolha = st.session_state.get(f"glmap_{_idx}", "— Ignorar —")
                            if _escolha and _escolha != "— Ignorar —" and _escolha.strip() in _linhas_ss:
                                _glmap_atual[_grupo] = _linhas_ss[_escolha.strip()]
                        rep = processar_fc(mp, _pos, _nf, _csv, _fat, _ped, _rat, _cart,
                                           out, mes_display,
                                           mapeamentos_extra=mapeamentos_usuario,
                                           grupo_linha_map=_glmap_atual)
                        with open(out,"rb") as f:
                            st.session_state["resultados"]["fc"] = (f.read(), rep, ext)
                    except Exception as e:
                        st.error(f"❌ Erro no FC Consolidado: {e}")
                        st.stop()

            if atualizar_espelho:
                with st.spinner("⚙️ Gerando Espelho DFC..."):
                    try:
                        import sys as _sys
                        _dfc_dir2 = os.path.dirname(os.path.abspath(__file__))
                        if _dfc_dir2 not in _sys.path: _sys.path.insert(0, _dfc_dir2)
                        import processar_espelho_dfc as _ped; import importlib; importlib.reload(_ped)
                        out_esp = os.path.join(tmpdir, "espelho_dfc.xlsx")
                        # Reuse already-saved files to avoid cursor issues
                        # FC: reuse if already saved by FC processing, else rewind and save
                        if paths.get("fc_saved") and os.path.exists(paths["fc_saved"]):
                            _fc_esp = paths["fc_saved"]
                        else:
                            fc_mestre.seek(0)
                            _fc_ext2 = "xlsm" if fc_mestre.name.endswith("xlsm") else "xlsx"
                            _fc_esp  = save_tmp(fc_mestre, f"fc_esp.{_fc_ext2}")
                        # CMV and estoque: rewind and save fresh
                        cmv_mestre.seek(0)
                        _cmv_esp = save_tmp(cmv_mestre,     "cmv_esp.xlsx")
                        estoque_mestre.seek(0)
                        _est_esp = save_tmp(estoque_mestre, "est_esp.xlsx")
                        # Usar bases do mês anterior para o espelho
                        _cart_esp = paths.get("cart_ant") or ([save_tmp(f, f"cart_esp_{i}.xlsx") for i,f in enumerate(cartoes_ant_files)] if cartoes_ant_files else [])
                        _rat_esp  = paths.get("rat_ant")  or (save_tmp(rateio_ant_file, "rat_esp.xlsx") if rateio_ant_file else None)
                        rep_esp = _ped.gerar_espelho_dfc(
                            cartoes_paths = _cart_esp,
                            rateio_path   = _rat_esp,
                            fc_path       = _fc_esp,
                            cmv_path      = _cmv_esp,
                            estoque_path  = _est_esp,
                            output_path   = out_esp,
                            mes_display   = mes_display,
                            ano           = 2026,
                        )
                        with open(out_esp, "rb") as f:
                            st.session_state["resultados"]["espelho"] = (f.read(), rep_esp)
                    except Exception as e:
                        st.error(f"❌ Erro no Espelho DFC: {e}")
                        st.stop()

            # ── Lançamento Manual: Vendas de Mercadorias ───────────────────────
            if lanc_vendas:
                with st.spinner("⚙️ Lançamento — Vendas de Mercadorias..."):
                    try:
                        import sys as _sys2
                        _dre_dir = os.path.dirname(os.path.abspath(__file__))
                        if _dre_dir not in _sys2.path: _sys2.path.insert(0, _dre_dir)
                        import processar_dre as _pdre; import importlib; importlib.reload(_pdre)
                        out_dre_v = os.path.join(tmpdir, "dre_vendas_out.xlsx")
                        _dre_v_path  = save_tmp(dre_vendas_mestre, "dre_vendas_m.xlsx")
                        _cmv_dre     = paths.get("cmv_saved") or save_tmp(cmv_mestre, "cmv_dre.xlsx")
                        _fat_dre     = paths.get("fat") or save_tmp(fat_file, "fat_dre.csv")
                        _ped_dre     = paths.get("ped") or save_tmp(pedidos_file, "ped_dre.xlsx")
                        rep_dre_v = _pdre.atualizar_vendas_mercadorias(
                            dre_path       = _dre_v_path,
                            fat_lojas_path = _fat_dre,
                            fat_vd_path    = _ped_dre,
                            cmv_path       = _cmv_dre,
                            mes_display    = mes_display,
                            ano            = 2026,
                            output_path    = out_dre_v,
                        )
                        with open(out_dre_v, "rb") as f:
                            st.session_state["resultados"]["dre_vendas"] = (f.read(), rep_dre_v)
                    except Exception as e:
                        st.error(f"❌ Erro em Vendas de Mercadorias: {e}")
                        st.stop()

            # ── Lançamento Manual: CMV ───────────────────────────────────────
            if lanc_cmv:
                with st.spinner("⚙️ Lançamento — CMV..."):
                    try:
                        import sys as _sys3
                        _dre_dir3 = os.path.dirname(os.path.abspath(__file__))
                        if _dre_dir3 not in _sys3.path: _sys3.path.insert(0, _dre_dir3)
                        import processar_dre as _pdre3; import importlib; importlib.reload(_pdre3)
                        out_dre_cmv   = os.path.join(tmpdir, "dre_cmv_out.xlsx")
                        _dre_cmv_path = save_tmp(dre_cmv_mestre, "dre_cmv_m.xlsx")
                        _cmv_csv      = paths.get("cmv_csv") or save_tmp(cmv_csv_file, "cmv_dre.csv")
                        rep_dre_cmv = _pdre3.atualizar_cmv(
                            dre_path     = _dre_cmv_path,
                            cmv_csv_path = _cmv_csv,
                            mes_display  = mes_display,
                            ano          = 2026,
                            output_path  = out_dre_cmv,
                        )
                        with open(out_dre_cmv, "rb") as f:
                            st.session_state["resultados"]["dre_cmv"] = (f.read(), rep_dre_cmv)
                    except Exception as e:
                        st.error(f"❌ Erro em CMV: {e}")
                        st.stop()

            if atualizar_dfc:
                with st.spinner("⚙️ DFC..."):
                    import sys
                    _dfc_dir = os.path.dirname(os.path.abspath(__file__))
                    if _dfc_dir not in sys.path:
                        sys.path.insert(0, _dfc_dir)
                    import processar_dfc as pdfc; importlib.reload(pdfc)
                    out = os.path.join(tmpdir, "dfc_out.xlsx")
                    mp  = save_tmp(dfc_mestre, "dfc_m.xlsb")

                    # FC: reset cursor before reading again (may have been read by FC processing)
                    fc_mestre.seek(0)
                    fc_ext    = "xlsm" if fc_mestre.name.lower().endswith("xlsm") else "xlsx"
                    fc_saved  = os.path.join(tmpdir, f"fc_for_dfc.{fc_ext}")
                    with open(fc_saved, "wb") as _f:
                        _f.write(fc_mestre.read())

                    # Estoque: reset cursor if already used, or skip
                    if estoque_mestre:
                        estoque_mestre.seek(0)
                        est_saved = os.path.join(tmpdir, "est_for_dfc.xlsx")
                        with open(est_saved, "wb") as _f:
                            _f.write(estoque_mestre.read())
                    else:
                        est_saved = None
                        st.warning("⚠️ DFC: Planilha de Estoque não carregada — col T não será preenchida.")

                    try:
                        rep = pdfc.processar_dfc(
                            dfc_path               = mp,
                            fc_consolidado_path    = fc_saved,
                            rateio_anterior_path   = paths["rat_ant"],
                            cartoes_anterior_paths = paths["cart_ant"],
                            estoque_mestre_path    = est_saved,
                            output_path            = out,
                            mes_display            = mes_display,
                            ano                    = 2026,
                        )
                        with open(out, "rb") as f:
                            st.session_state["resultados"]["dfc"] = (f.read(), rep)
                    except Exception as e:
                        st.error(f"❌ Erro no DFC: {e}")
                        st.stop()


                st.balloons()

    # ── Downloads e relatório — lidos do session_state, sobrevivem ao download ──
    resultados     = st.session_state.get("resultados", {})
    mes_processado = st.session_state.get("mes_processado", mes_display)

    if resultados:
        st.success(f"✅ **{mes_processado}** processado com sucesso!")

        st.markdown('<div class="section-title">⬇️ Downloads</div>', unsafe_allow_html=True)
        dc1, dc2, dc3 = st.columns(3)
        with dc1:
            if "cmv" in resultados:
                st.download_button(
                    "⬇️ CMV Atualizado",
                    data=resultados["cmv"][0],
                    file_name=f"CMV_VENA_2026_{mes_processado.upper()}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="dl_cmv",
                )
        with dc2:
            if "estoque" in resultados:
                st.download_button(
                    "⬇️ Estoque Atualizado",
                    data=resultados["estoque"][0],
                    file_name=f"ESTOQUE_VENA_2026_{mes_processado.upper()}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="dl_estoque",
                )
        with dc3:
            pass  # reservado

        if "espelho" in resultados:
                st.download_button(
                    "⬇️ Espelho DFC",
                    data=resultados["espelho"][0],
                    file_name=f"ESPELHO_DFC_{mes_processado.upper()}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="dl_espelho")
        dc4_e = st.columns(1)[0]
        with dc4_e:
            if "dfc" in resultados:
                st.download_button(
                    "⬇️ DFC Atualizado",
                    data=resultados["dfc"][0],
                    file_name=f"DFC_VENA_2026_{mes_processado.upper()}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="dl_dfc")
        if "dre_cmv" in resultados:
            st.download_button(
                "⬇️ DRE — CMV",
                data=resultados["dre_cmv"][0],
                file_name=f"DRE_CMV_{mes_processado.upper()}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="dl_dre_cmv",
            )

        if "dre_vendas" in resultados:
            st.download_button(
                "⬇️ DRE — Vendas de Mercadorias",
                data=resultados["dre_vendas"][0],
                file_name=f"DRE_VENDAS_{mes_processado.upper()}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="dl_dre_vendas",
            )

        dc4, dc5 = st.columns(2) if "fc" in resultados else (st.columns(1)[0], None)
        with dc4:
            if "fc" in resultados:
                st.download_button(
                    "⬇️ FC Consolidado Atualizado",
                    data=resultados["fc"][0],
                    file_name=f"FC_VENA_2026_{mes_processado.upper()}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="dl_fc",
                )

        st.markdown('<div class="section-title">📊 Relatório de Processamento</div>',
                    unsafe_allow_html=True)

        if "dre_cmv" in resultados:
            rep_dc = resultados["dre_cmv"][1]
            with st.expander("📉 DRE CMV — Detalhes"):
                c1, c2 = st.columns(2)
                c1.metric("✅ Lojas atualizadas", len(rep_dc.get("atualizadas", [])))
                c2.metric("⚠️ Não encontradas",  len(rep_dc.get("nao_encontradas", [])))
                if rep_dc.get("atualizadas"):
                    st.markdown("**Lançamentos realizados:**")
                    for item in rep_dc["atualizadas"]:
                        st.markdown(f"- {item}")
                if rep_dc.get("nao_encontradas"):
                    st.warning("Lojas sem valor no CMV Gerencial:")
                    for item in rep_dc["nao_encontradas"]:
                        st.markdown(f"- {item}")

        if "dre_vendas" in resultados:
            rep_dv = resultados["dre_vendas"][1]
            with st.expander("🛍️ DRE Vendas de Mercadorias — Detalhes"):
                c1, c2 = st.columns(2)
                c1.metric("✅ Lojas atualizadas", len(rep_dv.get("atualizadas", [])))
                c2.metric("⚠️ Não encontradas",  len(rep_dv.get("nao_encontradas", [])))
                if rep_dv.get("atualizadas"):
                    st.markdown("**Lançamentos realizados:**")
                    for item in rep_dv["atualizadas"]:
                        st.markdown(f"- {item}")
                if rep_dv.get("nao_encontradas"):
                    st.warning("Lojas sem valor nas bases:")
                    for item in rep_dv["nao_encontradas"]:
                        st.markdown(f"- {item}")

        if "cmv" in resultados:
            rep = resultados["cmv"][1]
            with st.expander("📊 CMV — Detalhes"):
                c1,c2,c3 = st.columns(3)
                c1.metric("Faturamento (LOJAS)", f"{len(rep['etapa1_ok'])} lojas")
                c2.metric("CMV/TRF inseridos",   f"{len(rep['etapa2_vd_ok'])+len(rep['etapa2_lojas_ok'])}")
                c3.metric("Pedidos VD",           f"{len(rep['etapa3_ok'])} lojas")
                if rep["etapa1_nok"]: st.warning(f"⚠️ Faturamento sem match: {rep['etapa1_nok']}")
                if rep["etapa2_nok"]: st.warning(f"⚠️ CMV sem match: {rep['etapa2_nok']}")
                if rep["etapa3_nok"]: st.warning(f"⚠️ Pedidos sem match: {rep['etapa3_nok']}")

        if "estoque" in resultados:
            rep = resultados["estoque"][1]
            with st.expander("📦 Estoque — Detalhes"):
                c1,c2,c3 = st.columns(3)
                c1.metric("Estoque gravado",  f"{len(rep['estoque_ok'])} lojas")
                c2.metric("Compras gravadas", f"{len(rep['compras_ok'])} lojas")
                c3.metric("CMV gravado",      f"{len(rep['cmv_ok'])} lojas")
                nok = list(set(rep["estoque_nok"]+rep["compras_nok"]+rep["cmv_nok"]))
                if nok: st.warning(f"⚠️ Lojas sem dados: {nok}")


        if "espelho" in resultados:
            rep = resultados["espelho"][1]
            with st.expander("📋 Espelho DFC — Detalhes"):
                st.write(f"**Aba 1 - Entradas:** {len(rep['aba1'].get('meses',[]))} meses")
                st.write(f"**Aba 2 - Compras:** {len(rep['aba2'].get('meses',[]))} meses")
                st.write(f"**Aba 3 - Parâmetros:** {len(rep['aba3'].get('lojas',[]))} lojas")
                st.write(f"**Aba 4 - Condições:** Cobert. Estoque = {rep['aba4'].get('prazo_medio',0)} dias")

        if "fc" in resultados:
            rep = resultados["fc"][1]
            with st.expander("💰 FC Consolidado — Detalhes"):
                st.write(f"**Coluna do mês:** {rep['mes_col']}")
                c1,c2,c3 = st.columns(3)
                c1.metric("Boletos (linha 11)", f"R$ {rep['f1_boletos']:,.2f}")
                c2.metric("Cartões (linha 12)", f"R$ {rep['f2_cartoes']:,.2f}")
                c3.metric("Estoque (linha 13)", f"R$ {rep['f3_estoque']:,.2f}")
                c1,c2,c3 = st.columns(3)
                c1.metric("Receita Líquida (L37)", f"R$ {rep['f5_receita']:,.2f}")
                c2.metric("CMV do Mês (L38)",      f"R$ {rep['f7_cmv']:,.2f}")
                c3.metric("Compras Mês (L39)",     f"R$ {rep['f6_compras']:,.2f}")
                if rep["f4_fornecedores"]:
                    st.write("**Fornecedores inseridos:**")
                    # Invert grupo_linha_usuario to get name from row number
                    _inv = {v:k for k,v in grupo_linha_usuario.items()} if grupo_linha_usuario else {}
                    for linha, val in rep["f4_fornecedores"].items():
                        nome = _inv.get(linha, f"Linha {linha}")
                        st.write(f"  Linha {linha} — {nome}: R$ {val:,.2f}")
                if rep.get("plano_contas_atualizado"):
                    st.success("✅ Plano de Contas F360 atualizado com " +
                               str(len(rep["plano_contas_atualizado"])) + " nova(s) conta(s):")
                    for item in rep["plano_contas_atualizado"]:
                        st.write(f"  ➕ {item}")
                if rep["f4_alertas"]:
                    st.markdown(
                        '<div class="alert-box">⚠️ <b>Contas ainda sem mapeamento:</b><br>'
                        + "<br>".join(f"• {a}" for a in rep["f4_alertas"])
                        + "</div>", unsafe_allow_html=True)

        if "dfc" in resultados:
            rep = resultados["dfc"][1]
            with st.expander("📈 DFC — Detalhes e Validação"):
                # Validação
                val = rep.get("validacao",{})
                st.write(f"**Mês anterior para comparação:** {val.get('mes_anterior','')}")
                st.markdown("**🔍 Comparativo Cartões (base vs FC mês anterior):**")
                vc = val.get('cartoes',{})
                cc1,cc2,cc3 = st.columns(3)
                cc1.metric("Base Crédito", f"R$ {vc.get('base_credito',0):,.2f}")
                cc2.metric("FC Cartões ant.", f"R$ {vc.get('fc_cartoes_ant',0):,.2f}")
                cc3.metric("Diferença", f"R$ {vc.get('diff_credito',0):,.2f}",
                           delta_color="off" if abs(vc.get('diff_credito',0)) < 1 else "inverse")
                cb1,cb2,cb3 = st.columns(3)
                cb1.metric("Base Boleto", f"R$ {vc.get('base_boleto',0):,.2f}")
                cb2.metric("FC Boletos ant.", f"R$ {vc.get('fc_boletos_ant',0):,.2f}")
                cb3.metric("Diferença", f"R$ {vc.get('diff_boleto',0):,.2f}",
                           delta_color="off" if abs(vc.get('diff_boleto',0)) < 1 else "inverse")
                st.markdown("**🔍 Comparativo Compras (base vs FC mês anterior):**")
                vcomp = val.get('compras',{})
                cm1,cm2,cm3 = st.columns(3)
                cm1.metric("Base Mercadorias", f"R$ {vcomp.get('base_mercadorias',0):,.2f}")
                cm2.metric("FC Mercadorias ant.", f"R$ {vcomp.get('fc_mercadorias_ant',0):,.2f}")
                cm3.metric("Diferença", f"R$ {vcomp.get('diff',0):,.2f}",
                           delta_color="off" if abs(vcomp.get('diff',0)) < 1 else "inverse")
                # Abas atualizadas
                cc = rep.get('condicoes_compra',{})
                pc = rep.get('projecao_compra',{})
                pp = rep.get('parametros_proj',{})
                pe = rep.get('projecao_entradas',{})
                st.write(f"**Condições de Compra:** Prazo Médio Estoque = {cc.get('valor',0):.1f} dias → {len(cc.get('lojas',[]))} lojas")
                st.write(f"**Projeção Compra:** L173={pc.get('L173_mercadorias',0):,.2f} | L174={pc.get('L174_trf',0):,.2f} | L175={pc.get('L175_mkt',0):,.2f}")
                st.write(f"**Projeção Entradas:** Crédito={pe.get('L14_credito',0):,.2f} | Boleto={pe.get('L15_boleto',0):,.2f}")
                if pp.get('sem_estoque'):
                    st.warning(f"⚠️ Lojas sem estoque no mês anterior: {pp['sem_estoque']}")


if __name__ == "__main__":
    import sys
    if "streamlit" in sys.modules or any("streamlit" in a for a in sys.argv):
        run_streamlit()
    else:
        print("Execute com: python -m streamlit run update_cmv_vena.py")
